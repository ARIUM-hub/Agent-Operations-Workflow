from __future__ import annotations

import csv
import math
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal
from io import BytesIO, StringIO
from pathlib import Path

from openpyxl import load_workbook

BATCH_LIMIT = 50
BATCH_TEXT_KEYS = (
    "conversation",
    "conversation_text",
    "message",
    "text",
    "content",
    "body",
    "客服会话",
    "会话",
    "消息",
    "内容",
)
STORE_KEYS = ("store_name", "store", "shop_name", "shop", "店铺名称", "店铺")
SKU_KEYS = ("sku", "seller_sku", "merchant_sku", "product_sku", "商品sku", "商品 sku")
PRODUCT_ID_KEYS = (
    "platform_product_id",
    "product_id",
    "asin",
    "item_id",
    "listing_id",
    "商品id",
    "商品 id",
)


@dataclass(frozen=True)
class ExtractedConversation:
    conversation_text: str
    store_name: str = ""
    sku: str = ""
    platform_product_id: str = ""


def extract_conversation(filename: str, content: bytes) -> ExtractedConversation:
    suffix = Path(filename).suffix.lower()
    if suffix in {".txt", ".log"}:
        extracted = ExtractedConversation(conversation_text=_decode_text(content))
    elif suffix == ".csv":
        extracted = _extract_single_csv(content)
    elif suffix in {".xlsx", ".xlsm"}:
        extracted = _extract_single_xlsx(content)
    else:
        raise ValueError("仅支持 txt、log、csv、xlsx 或 xlsm 文件")
    return _clean_extracted(extracted, "没有可分析内容")


def extract_batch_conversations(
    filename: str,
    content: bytes,
    limit: int = BATCH_LIMIT,
) -> list[ExtractedConversation]:
    suffix = Path(filename).suffix.lower()
    if suffix in {".txt", ".log"}:
        items = [
            ExtractedConversation(conversation_text=value)
            for value in _split_batch_text(_decode_text(content))
        ]
    elif suffix == ".csv":
        items = _extract_batch_csv_rows(content)
    elif suffix in {".xlsx", ".xlsm"}:
        items = _extract_batch_xlsx_rows(content)
    else:
        raise ValueError("仅支持 txt、log、csv、xlsx 或 xlsm 文件")
    cleaned = [
        _clean_extracted(item, "文件中没有可分析会话，请检查导出内容")
        for item in items
        if item.conversation_text and item.conversation_text.strip()
    ]
    if not cleaned:
        raise ValueError("文件中没有可分析会话，请检查导出内容")
    if len(cleaned) > limit:
        raise ValueError(f"单次最多分析 {limit} 条，请拆分文件后重试")
    return cleaned


def _decode_text(content: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-8", "gb18030"):
        try:
            return content.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise ValueError("文件编码无法识别，请使用 UTF-8 导出")


def _extract_csv(content: bytes) -> str:
    decoded = _decode_text(content)
    reader = csv.DictReader(StringIO(decoded))
    rows: list[str] = []
    for row in reader:
        speaker = _first_value(
            row,
            ("speaker", "role", "sender", "from", "角色", "发送方"),
        ) or "Unknown"
        message = _first_value(row, ("message", "text", "content", "body", "消息", "内容"))
        if message:
            rows.append(f"{speaker}: {message}")
    if rows:
        return "\n".join(rows)
    return decoded


def _split_batch_text(text: str) -> list[str]:
    normalized = text.replace("\r\n", "\n")
    for separator in ("\n---\n", "\n===\n"):
        if separator in normalized:
            return normalized.split(separator)
    return [normalized]


def _extract_xlsx(content: bytes) -> str:
    workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    sheet = workbook.active
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        return ""

    headers = [str(value).strip() if value is not None else "" for value in rows[0]]
    speaker_index = _first_index(
        headers,
        ("speaker", "role", "sender", "from", "角色", "发送方"),
    )
    message_index = _first_index(
        headers,
        ("message", "text", "content", "body", "消息", "内容"),
    )
    lines: list[str] = []

    for row in rows[1:]:
        values = ["" if value is None else str(value).strip() for value in row]
        if message_index is not None and message_index < len(values):
            speaker = (
                values[speaker_index]
                if speaker_index is not None and speaker_index < len(values)
                else "Unknown"
            )
            message = values[message_index]
            if message:
                lines.append(f"{speaker}: {message}")
        else:
            line = " ".join(value for value in values if value)
            if line:
                lines.append(line)
    return "\n".join(lines)


def _metadata_from_mapping(row: dict[str, object]) -> dict[str, str]:
    normalized = {
        str(key).strip().casefold(): value
        for key, value in row.items()
        if key is not None
    }
    return {
        "store_name": _first_identifier(normalized, STORE_KEYS),
        "sku": _first_identifier(normalized, SKU_KEYS),
        "platform_product_id": _first_identifier(normalized, PRODUCT_ID_KEYS),
    }


def _first_identifier(row: dict[str, object], keys: tuple[str, ...]) -> str:
    for key in keys:
        value = row.get(key.casefold())
        if value is None or value == "":
            continue
        text = _identifier_text(value)
        if text:
            return text
    return ""


def _identifier_text(value: object) -> str:
    if isinstance(value, bool) or isinstance(value, (date, datetime)):
        raise ValueError("商品标识必须是文本或数字")
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        if not math.isfinite(value):
            raise ValueError("商品标识必须是有限数字")
        if value.is_integer():
            return str(int(value))
        return format(Decimal(str(value)), "f").rstrip("0").rstrip(".")
    if not isinstance(value, str):
        raise ValueError("商品标识必须是文本或数字")
    return value.strip()


def _consistent_metadata(rows: list[dict[str, object]]) -> dict[str, str]:
    metadata_rows = [_metadata_from_mapping(row) for row in rows]
    result = {"store_name": "", "sku": "", "platform_product_id": ""}
    labels = {
        "store_name": "店铺",
        "sku": "SKU",
        "platform_product_id": "平台商品 ID",
    }
    for field, label in labels.items():
        values = [row[field] for row in metadata_rows if row[field]]
        normalized = {value.casefold() for value in values}
        if len(normalized) > 1:
            raise ValueError(f"单条文件包含多个 {label}，请改用批量上传")
        if values:
            result[field] = values[0]
    return result


def _clean_extracted(
    item: ExtractedConversation,
    empty_message: str,
) -> ExtractedConversation:
    conversation_text = item.conversation_text.strip()
    if not conversation_text:
        raise ValueError(empty_message)
    return ExtractedConversation(
        conversation_text=conversation_text,
        store_name=item.store_name.strip(),
        sku=item.sku.strip(),
        platform_product_id=item.platform_product_id.strip(),
    )


def _csv_mapping_rows(content: bytes) -> tuple[str, list[dict[str, object]]]:
    decoded = _decode_text(content)
    rows = [dict(row) for row in csv.DictReader(StringIO(decoded))]
    return decoded, rows


def _xlsx_mapping_rows(content: bytes) -> list[dict[str, object]]:
    workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    sheet = workbook.active
    values = list(sheet.iter_rows(values_only=True))
    if not values:
        return []
    headers = ["" if value is None else str(value).strip() for value in values[0]]
    return [
        {
            headers[index]: value
            for index, value in enumerate(row)
            if index < len(headers) and headers[index]
        }
        for row in values[1:]
    ]


def _conversation_value(row: dict[str, object]) -> str:
    normalized = {
        str(key).strip().casefold(): value
        for key, value in row.items()
        if key is not None
    }
    for key in BATCH_TEXT_KEYS:
        value = normalized.get(key.casefold())
        if value is not None and str(value).strip():
            return str(value).strip()
    return ""


def _extract_single_csv(content: bytes) -> ExtractedConversation:
    _, rows = _csv_mapping_rows(content)
    metadata = _consistent_metadata(rows)
    return ExtractedConversation(
        conversation_text=_extract_csv(content),
        **metadata,
    )


def _extract_single_xlsx(content: bytes) -> ExtractedConversation:
    rows = _xlsx_mapping_rows(content)
    metadata = _consistent_metadata(rows)
    return ExtractedConversation(
        conversation_text=_extract_xlsx(content),
        **metadata,
    )


def _extract_batch_csv_rows(content: bytes) -> list[ExtractedConversation]:
    decoded, rows = _csv_mapping_rows(content)
    extracted = [
        ExtractedConversation(
            conversation_text=_conversation_value(row),
            **_metadata_from_mapping(row),
        )
        for row in rows
        if _conversation_value(row)
    ]
    if extracted:
        return extracted
    return [
        ExtractedConversation(conversation_text=value)
        for value in _split_batch_text(decoded)
    ]


def _extract_batch_xlsx_rows(content: bytes) -> list[ExtractedConversation]:
    rows = _xlsx_mapping_rows(content)
    conversation_keys = {key.casefold() for key in BATCH_TEXT_KEYS}
    has_conversation_column = bool(rows) and any(
        str(key).strip().casefold() in conversation_keys for key in rows[0]
    )
    return [
        ExtractedConversation(
            conversation_text=(
                _conversation_value(row)
                if has_conversation_column
                else " ".join(
                    str(value).strip()
                    for value in row.values()
                    if value is not None and str(value).strip()
                )
            ),
            **_metadata_from_mapping(row),
        )
        for row in rows
        if not has_conversation_column or _conversation_value(row)
    ]


def _first_value(
    row: dict[str, str | None],
    keys: tuple[str, ...],
) -> str | None:
    normalized = {key.strip().lower(): value for key, value in row.items() if key}
    for key in keys:
        value = normalized.get(key.lower())
        if value and value.strip():
            return value.strip()
    return None


def _first_index(headers: list[str], keys: tuple[str, ...]) -> int | None:
    lowered = [header.lower() for header in headers]
    for key in keys:
        if key.lower() in lowered:
            return lowered.index(key.lower())
    return None
